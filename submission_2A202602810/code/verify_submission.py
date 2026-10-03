"""Check submission evidence, workbook formulas and real notebook execution."""
import ast
import json
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
import nbformat

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent if (HERE.parent / 'scripts').exists() else HERE.parent.parent
OUT = ROOT / 'submission_2A202602810'

def main():
    required = ['REPORT.md','experiments.xlsx','predictions_eval.csv','eval_result.json','code/lab.ipynb']
    for name in required:
        assert (OUT/name).is_file(), name
    xlsx = openpyxl.load_workbook(OUT/'experiments.xlsx',data_only=True)
    formulas = openpyxl.load_workbook(OUT/'experiments.xlsx',data_only=False)
    template_path = ROOT/'templates/experiment_table_template.xlsx'
    if not template_path.exists():
        template_path = HERE/'experiment_table_template.xlsx'
    template = openpyxl.load_workbook(template_path)
    assert formulas.sheetnames == template.sheetnames
    assert [c.value for c in formulas['Experiments'][1]] == [c.value for c in template['Experiments'][1]]
    # Existing template formulas remain intact and every calculated cache is error-free.
    for sheet in template:
        for row in sheet:
            for c in row:
                if c.data_type=='f':
                    assert formulas[sheet.title][c.coordinate].value==c.value,(sheet.title,c.coordinate)
    for sheet in xlsx:
        assert not any(c.data_type=='e' for row in sheet for c in row),sheet.title
    assert xlsx['Summary']['D13'].value==7
    headers=[c.value for c in xlsx['Experiments'][1]]
    rows=[dict(zip(headers,[c.value for c in row])) for row in xlsx['Experiments'].iter_rows(min_row=2) if row[0].value]
    ids=[r['exp_id'] for r in rows]
    assert len(ids)==len(set(ids))
    logs={p.stem:json.loads(p.read_text()) for p in (OUT/'results').glob('*.json')}
    assert set(ids)==set(logs)
    for row in rows:
        result=logs[row['exp_id']]
        assert (OUT/row['figure_file']).is_file()
        assert result['summary']['completed_epochs']==row['epochs']==20
        assert len(result['history']['epoch'])==20
        assert len(result['grad_norm_steps'])==sum(result['history']['steps'])
        assert sum(v is None for v in result['grad_norm_steps'])==result['summary'].get('scaler_skipped_steps',0)
        assert all(s==int(np.ceil(371847/result['cfg']['batch'])) for s in result['history']['steps'])
        assert all(np.isfinite(v) for key in ('train_loss','val_loss','grad_norm','val_macro_f1') for v in result['history'][key])
        for key in ('step0_loss','best_val_loss','val_acc','val_macro_f1','final_train_loss','final_val_loss','time_per_epoch_s','peak_mem_MB'):
            assert abs(row[key]-result['summary'][key])<1e-12,(row['exp_id'],key)
    baseline=[r for r in rows if r['group']=='baseline']
    assert len(baseline)==3
    measured=np.array([r['val_macro_f1'] for r in baseline])
    assert abs(xlsx['Seeds']['C8'].value-measured.mean())<1e-12
    assert abs(xlsx['Seeds']['C10'].value-2*measured.std(ddof=1))<1e-12
    selection=json.loads((OUT/'selection.json').read_text())
    eligible=[r for r in rows if r['seed']==1 and r['exp_id']!='amp-fp16' and not r['diverged']]
    assert selection['final_exp_id']==max(eligible,key=lambda r:r['val_macro_f1'])['exp_id']
    for row in rows:
        if row['exp_id'] not in (selection['baseline_exp_id'],selection['final_exp_id']):
            assert row['eval_acc'] is None and row['eval_macro_f1'] is None
    official=json.loads((OUT/'eval_result.json').read_text())
    pred=pd.read_csv(OUT/'predictions_eval.csv')
    assert list(pred.columns)==['row_id','pred'] and len(pred)==116203
    assert pred.row_id.is_unique and pred.pred.between(0,6).all()
    with np.load(ROOT/'data/processed/eval.npz') as a:
        assert set(pred.row_id)==set(a['row_id'].tolist())
        aligned=pred.set_index('row_id').loc[a['row_id'],'pred'].to_numpy()
        cm=np.zeros((7,7),dtype=np.int64)
        np.add.at(cm,(a['y'],aligned),1)
    assert np.array_equal(cm,official['confusion_matrix'])
    final_row=next(r for r in rows if r['exp_id']==selection['final_exp_id'])
    assert final_row['eval_acc']==official['accuracy'] and final_row['eval_macro_f1']==official['macro_f1']
    report=(OUT/'REPORT.md').read_text()
    assert f"{official['accuracy']:.6f}" in report and f"{official['macro_f1']:.6f}" in report
    notebook=nbformat.read(OUT/'code/lab.ipynb',as_version=4)
    code_cells=[c for c in notebook.cells if c.cell_type=='code']
    assert all(c.execution_count is not None for c in code_cells)
    assert not any(o.output_type=='error' for c in code_cells for o in c.outputs)
    for p in (OUT/'code').glob('*.py'):
        tree=ast.parse(p.read_text())
        for node in ast.walk(tree):
            if isinstance(node,ast.Raise) and node.exc is not None:
                exc=node.exc.func if isinstance(node.exc,ast.Call) else node.exc
                assert not isinstance(exc,ast.Name) or exc.id!='Not'+'Implemented'+'Error',p
    assert not list(OUT.rglob('*.pt')) and not list(OUT.rglob('*.npz'))
    print(f'PASS: {len(rows)} experiments, 20 epochs each, 7 topics, 3 baseline seeds.')
    print('PASS: figures, original Excel headers/formulas, cached values, JSON and report match.')
    print(f'PASS: {len(code_cells)} executed notebook cells; {len(pred)} valid eval predictions.')
    print(f"Final: {selection['final_exp_id']}; accuracy={official['accuracy']:.6f}; macro-F1={official['macro_f1']:.6f}")

if __name__=='__main__':
    main()
