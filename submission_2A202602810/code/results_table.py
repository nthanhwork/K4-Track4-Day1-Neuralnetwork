"""Populate the supplied workbook without changing sheets, headers or formulas."""
import json
import math
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
import numpy as np
import openpyxl

FORMULA_COLUMNS = {'step0_gap_vs_lnC', 'gap_val_minus_train', 'delta_val_f1_vs_base', 'beyond_noise'}

def save_result(result, results_dir='../results'):
    path = Path(results_dir) / (result['cfg']['exp_id'] + '.json')
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {k: v for k, v in result.items() if k != 'best_state'}
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, allow_nan=False))
    return str(path)

def load_results(results_dir='../results'):
    return [json.loads(p.read_text()) for p in sorted(Path(results_dir).glob('*.json'))]

def to_row(result, eval_scores=None, notes=''):
    cfg = result['cfg']
    row = {**cfg, **result['summary'], 'hidden': '-'.join(map(str, cfg['hidden'])),
           'clip_norm': cfg['clip_norm'] if cfg['clip_norm'] is not None else 'none',
           'figure_file': f"figures/{cfg['exp_id']}.png"}
    row['notes'] = '; '.join(filter(None, [cfg.get('prediction', ''), cfg.get('notes', ''),
                          f"memory={result['summary']['memory_kind']}; train loss: fixed 50000 probe, eval mode", notes]))
    if eval_scores is not None:
        row.update(eval_acc=eval_scores['accuracy'], eval_macro_f1=eval_scores['macro_f1'])
    return row

def _cache_formulas(path, values):
    """Cache measured formula values for readers without Excel; retain every formula.

    Standard XLSX stores a formula and its last calculated value separately.
    Values here are computed from the same experiment rows, never invented.
    """
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with ZipFile(path) as z:
        members = [(item, z.read(item.filename)) for item in z.infolist()]
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for item, content in members:
            if item.filename in values:
                root = ET.fromstring(content)
                for c in root.iter(ns + 'c'):
                    if c.find(ns + 'f') is None:
                        continue
                    value = values[item.filename].get(c.attrib['r'], '')
                    v = c.find(ns + 'v')
                    if v is None:
                        v = ET.SubElement(c, ns + 'v')
                    if isinstance(value, str):
                        c.set('t', 'str'); v.text = value
                    else:
                        c.attrib.pop('t', None)
                        v.text = str(value)
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            z.writestr(item, content)

def write_xlsx(rows, template_path, out_path, annotations=None):
    if len(rows) > 60:
        raise ValueError('Template supports 60 experiment rows')
    wb = openpyxl.load_workbook(template_path)
    assert wb.sheetnames == ['Legend', 'Experiments', 'Seeds', 'Summary']
    ws = wb['Experiments']
    headers = [c.value for c in ws[1]]
    for rr in range(2, 62):
        for cc, key in enumerate(headers, 1):
            if key not in FORMULA_COLUMNS:
                ws.cell(rr, cc).value = None
    for rr, row in enumerate(rows, 2):
        for cc, key in enumerate(headers, 1):
            if key not in FORMULA_COLUMNS:
                ws.cell(rr, cc).value = row.get(key)
    baseline = [r for r in rows if r['group'] == 'baseline']
    assert 2 <= len(baseline) <= 5
    seeds = wb['Seeds']
    for rr in range(2, 7):
        seeds.cell(rr, 1).value = baseline[rr-2]['exp_id'] if rr-2 < len(baseline) else None
    mean = float(np.mean([r['val_macro_f1'] for r in baseline]))
    noise = 2 * float(np.std([r['val_macro_f1'] for r in baseline], ddof=1))
    values = {'xl/worksheets/sheet2.xml': {}, 'xl/worksheets/sheet3.xml': {}, 'xl/worksheets/sheet4.xml': {}}
    for rr, row in enumerate(rows, 2):
        cache = values['xl/worksheets/sheet2.xml']
        cache[f'AD{rr}'] = row['step0_loss'] - math.log(7)
        cache[f'AE{rr}'] = row['final_val_loss'] - row['final_train_loss'] if row['final_val_loss'] is not None else ''
        cache[f'AF{rr}'] = row['val_macro_f1'] - mean if row['val_macro_f1'] is not None else ''
        cache[f'AG{rr}'] = ('Có' if abs(row['val_macro_f1'] - mean) > noise else 'Không') if row['val_macro_f1'] is not None else ''
    for col, key in [('B', 'val_acc'), ('C', 'val_macro_f1'), ('D', 'best_val_loss')]:
        cache = values['xl/worksheets/sheet3.xml']
        observations = [r[key] for r in baseline]
        for rr, value in enumerate(observations, 2):
            cache[f'{col}{rr}'] = value
        cache[f'{col}8'] = float(np.mean(observations))
        cache[f'{col}9'] = float(np.std(observations, ddof=1))
        cache[f'{col}10'] = 2 * cache[f'{col}9']
    summary = wb['Summary']
    for rr in range(2, 12):
        group = summary.cell(rr, 1).value
        grouped = [r for r in rows if r['group'] == group]
        valid = [r for r in grouped if r['val_macro_f1'] is not None]
        cache = values['xl/worksheets/sheet4.xml']
        for col, value in [('B', len(grouped)), ('C', len(valid)),
                           ('D', max((r['val_macro_f1'] for r in valid), default='')),
                           ('E', min((r['val_macro_f1'] for r in valid), default='')),
                           ('F', max((r['val_acc'] for r in valid), default=''))]:
            cache[f'{col}{rr}'] = value
        if group in ('loss', 'optimizer', 'hparam', 'dropout', 'clipping', 'amp', 'init'):
            cache[f'G{rr}'] = 'Có' if valid else 'Chưa'
        if valid:
            best = max(valid, key=lambda r: r['val_macro_f1'])
            summary.cell(rr, 8).value = f"Best: {best['exp_id']}, F1={best['val_macro_f1']:.6f}; compare against 2σ={noise:.6f}."
    values['xl/worksheets/sheet4.xml']['D13'] = sum(any(r['group'] == g and r['val_macro_f1'] is not None for r in rows)
                                                    for g in ('loss','optimizer','hparam','dropout','clipping','amp','init'))
    if annotations:
        legend = wb['Legend']
        legend.cell(59, 1).value = 'Measured supplementary evidence (associated exp_id in column A)'
        for rr, row in enumerate(annotations, 60):
            for cc, value in enumerate(row, 1):
                legend.cell(rr, cc).value = value
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    _cache_formulas(out_path, values)
